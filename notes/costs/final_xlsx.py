# /// script
# dependencies = ["openpyxl"]
# ///
import json, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
D = "model-costs"
OUT = "notes/costs/painting_models_final_2026-09-28.xlsx"
ms = {m["id"]: m for m in json.load(open(f"{D}/or_models.json"))["data"]}
sets = json.load(open(f"{D}/sets.json"))
latest = re.findall(r"`([^`]+)`", open("notes/costs/latest_models.md").read().split("| # |")[1])
latest = [i for i in latest if i != "perceptron/perceptron-mk1.5"]
def f(x):
    try: return float(x)
    except (TypeError, ValueError): return None
OURS = {"anthropic/claude-opus-5.5", "anthropic/claude-fable-5.1", "openai/gpt-6-astra", "google/gemini-3.8-flash", "openai/gpt-6-luna",
        "moonshotai/kimi-k3", "xiaomi/mimo-v2.6-pro", "meta/muse-spark-1.3", "z-ai/glm-5.3-flash", "deepseek/deepseek-v4.1-flash", "stealth/space-bunny-alpha"}
NOTE = {"unbiased/pareto": "ensemble (several unnamed models per request): kept on purpose",
        "mistralai/mistral-medium-3-5": "cached tokens at 10% per Mistral (not listed on OpenRouter)",
        "fireworks/ember-1": "Kimi K3 fine-tune with shorter reasoning (likely cheaper than shown)",
        "google/gemini-3.5-flash-lite": "Flash Lite line (3.5 is its newest)", "stealth/space-bunny-alpha": "free (stealth)",
        "x-ai/grok-build-0.1": "xAI's coding model", "meta/muse-spark-1.3-contributor": "cheap contributor tier (check its data terms)"}
CACHE_OVERRIDE = {"mistralai/mistral-medium-3-5": 0.15}
def reason(i):
    if i.startswith("sakana/fugu"): return "orchestrator (routes between models)"
    if i == "sakana/sakana-namazu": return "Japanese fine-tune of Kimi K2.6"
    if re.match(r"openai/(.*-pro|o\d-pro)$", i): return "OpenAI Pro: thinks far longer, costs far more than its list price suggests"
    if "-image" in i: return "image-generation model"
    if i in ("rekaai/reka-edge", "perceptron/perceptron-mk1.5", "z-ai/glm-4.5v"):
        return "context too small for a painting session (16K / 36K / 65K)" + (": a robotics model" if "perceptron" in i else "")
    if i == "openai/o1": return "Pro-like: built to think far longer (about $1,049 a pass at list price)"
    DUP = {"google/gemini-3.1-flash-lite-preview": "google/gemini-3.1-flash-lite", "google/gemini-3.1-pro-preview-customtools": "google/gemini-3.1-pro-preview",
           "google/gemini-2.5-pro-preview": "google/gemini-2.5-pro", "openai/gpt-4o": "openai/gpt-4o-2024-05-13 and -2024-11-20 (first and last GPT-4o kept)",
           "openai/gpt-4o-2024-08-06": "openai/gpt-4o-2024-05-13 and -2024-11-20 (first and last GPT-4o kept)",
           "openai/gpt-4o-mini-2024-07-18": "openai/gpt-4o-mini", "qwen/qwen3.5-plus-02-15": "qwen/qwen3.5-plus-20260420"}
    if i in DUP: return "duplicate of " + DUP[i]
    return None
allp = [i for i in sets["all"] if not reason(i)]
excluded = sorted({i: reason(i) for i in sets["all"] if reason(i)}.items())
excluded = [(i, w) for i, w in excluded]
excluded += [("openai/gpt-5.5", "superseded by GPT-6 (latest list only)"), ("openai/gpt-chat-latest", "ChatGPT's rolling alias (latest list only)"),
             ("xiaomi/mimo-v2.5", "older MiMo (latest list only)")]
NOTE.update({"openai/gpt-4-turbo": "historical: kept for fun (no cache price)", "anthropic/claude-opus-4.1": "historical Opus (old pricing)",
             "openai/o4-mini-high": "o4-mini with reasoning effort high (kept separate)", "openai/gpt-4o-2024-05-13": "first GPT-4o",
             "openai/gpt-4o-2024-11-20": "last GPT-4o"})
bold = Font(bold=True); hdr = PatternFill("solid", fgColor="DDD6C8"); ours = PatternFill("solid", fgColor="F3E6B8"); blue = Font(color="1F4E9A")
wb = Workbook()
# Profiles
pr = wb.active; pr.title = "Profiles"
pr.append(["Token profile of one painting in 4 sittings (edit the blue cells: every cost recalculates)"]); pr["A1"].font = bold
pr.append(["profile", "prompt tokens", "output tokens", "images", "based on"])
for c in pr[2]: c.font = bold; c.fill = hdr
pr.append(["light", 4_600_000, 72_000, 55, "r19 Luna F (stopped early, 36 chunks)"])
pr.append(["typical", 40_000_000, 220_000, 200, "median of our 9 full 4-sitting paintings (rounds 17-19)"])
pr.append(["heavy", 75_000_000, 600_000, 300, "r19 Gemini B, r18 Space Bunny"])
pr.append([]); pr.append(["prompt split", "share"]); pr["A7"].font = bold
pr.append(["cache read", 0.89]); pr.append(["uncached input", 0.08]); pr.append(["cache write", 0.03])
for r in (3, 4, 5):
    for c in "BCD": pr[f"{c}{r}"].font = blue; pr[f"{c}{r}"].number_format = "#,##0"
for r in (8, 9, 10): pr[f"B{r}"].font = blue; pr[f"B{r}"].number_format = "0%"
for col, w in zip("ABCDE", (16, 16, 14, 9, 55)): pr.column_dimensions[col].width = w
ROW = {"light": 3, "typical": 4, "heavy": 5}
COLS = ["#", "model", "released", "note", "painted with it", "$/M input", "$/M output", "$/M cache read", "$/M cache write", "$/image",
        "light $", "typical $", "heavy $", "one pass $", "three passes $"]
def sheet(title, ids):
    ws = wb.create_sheet(title); ws.append(COLS)
    for c in ws[1]: c.font = bold; c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="top")
    for n, i in enumerate(ids, start=2):
        m = ms[i]; p = m["pricing"]; pm = lambda k: f(p.get(k)) * 1e6 if f(p.get(k)) is not None else None
        cr = CACHE_OVERRIDE.get(i, pm("input_cache_read"))
        ws.append([n - 1, i, datetime.date.fromtimestamp(m["created"]).isoformat(), NOTE.get(i, ""), "yes" if i in OURS else "",
                   pm("prompt"), pm("completion"), cr, pm("input_cache_write"), f(p.get("image")) if f(p.get("image")) else None])
        def cost(prof):
            r = ROW[prof]
            return (f'=(Profiles!$B${r}*(Profiles!$B$9*F{n}+Profiles!$B$8*IF(H{n}="",F{n},H{n})+Profiles!$B$10*IF(I{n}="",F{n},I{n}))'
                    f'+Profiles!$C${r}*G{n})/1000000+Profiles!$D${r}*IF(J{n}="",0,J{n})')
        ws[f"K{n}"] = cost("light"); ws[f"L{n}"] = cost("typical"); ws[f"M{n}"] = cost("heavy")
        ws[f"N{n}"] = f"=K{n}+L{n}+M{n}"; ws[f"O{n}"] = f"=3*N{n}"
        for c in "FGHIJ": ws[f"{c}{n}"].number_format = "0.000"
        for c in "KLMNO": ws[f"{c}{n}"].number_format = "$#,##0.00"
        if i in OURS:
            for c in ws[n]: c.fill = ours
    last = len(ids) + 1; t = last + 2
    ws[f"B{t}"] = "total"; ws[f"B{t}"].font = bold
    for c in "KLMNO":
        ws[f"{c}{t}"] = f"=SUM({c}2:{c}{last})"; ws[f"{c}{t}"].number_format = "$#,##0"; ws[f"{c}{t}"].font = bold
    for k, w in enumerate((5, 40, 11, 44, 9, 9, 9, 9, 9, 8, 10, 10, 10, 11, 12), 1): ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:O{last}"
    return t
tl = sheet(f"Latest ({len(latest)})", latest)
ta = sheet(f"Present & past ({len(allp)})", allp)
# Excluded
ex = wb.create_sheet("Excluded"); ex.append(["model", "why"])
for c in ex[1]: c.font = bold; c.fill = hdr
for i, why in excluded: ex.append([i, why])
ex.column_dimensions["A"].width = 44; ex.column_dimensions["B"].width = 80
# Summary
su = wb.create_sheet("Summary", 0)
L = f"'Latest ({len(latest)})'"; A = f"'Present & past ({len(allp)})'"
su.append(["Vision models that could paint, priced on OpenRouter (2026-09-28)"]); su["A1"].font = Font(bold=True, size=13)
su.append([])
su.append(["set", "models", "light $", "typical $", "heavy $", "one pass $", "three passes $"])
for c in su[3]: c.font = bold; c.fill = hdr
su.append(["Latest", len(latest)] + [f"={L}!{c}{tl}" for c in "KLMNO"])
su.append(["Present & past", len(allp)] + [f"={A}!{c}{ta}" for c in "KLMNO"])
su.append(["Sum of both", "=B4+B5"] + [f"={c}4+{c}5" for c in "CDEFG"])
for r in (4, 5, 6):
    for c in "CDEFG": su[f"{c}{r}"].number_format = "$#,##0"
su["A6"].font = bold
notes = ["",
 "One pass = one light, one typical and one heavy painting per model (4 sittings each; profiles on the Profiles sheet).",
 "Latest: the newest model per line released since 2026-03-28. Present & past: every version still on OpenRouter.",
 "Both: images in, tool calls, interactive; no routers or orchestrators (Unbiased's Pareto, an ensemble, kept on purpose);",
 "no OpenAI Pro models or o1; no image-generation models; no duplicates (o4-mini and -high kept apart; first and last GPT-4o);",
 "no model whose context can't hold a session (under 128K). The sum counts the latest models twice (they're in both sets).",
 "Cost = prompt x (8% input + 89% cache read + 3% cache write) + output + images, at list prices; no cache price = input price.",
 "Ballpark: x0.5 to x2 (tokenizers, image tokens and thinking length differ). If caching fails, x3 to x5.",
 "Source: https://openrouter.ai/api/v1/models (fetched 2026-09-28); token profiles from our session logs.",
 "Yellow rows: models we have painted with."]
for line in notes: su.append([line])
for k, w in enumerate((18, 9, 11, 11, 11, 12, 14), 1): su.column_dimensions[get_column_letter(k)].width = w
wb.move_sheet("Profiles", offset=len(wb.sheetnames))
wb.save(OUT); print(OUT, "| latest", len(latest), "| present & past", len(allp), "| excluded", len(excluded))
